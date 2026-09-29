"""Explicit Excel-to-SQLite catalog seeding; not used by runtime tools."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from backend.database.db import CATALOG_COLUMNS, DEFAULT_DATABASE_PATH, SearchDatabase

DEFAULT_DATASET_PATH = Path(__file__).resolve().parents[2] / "data" / "TELORA_Synthetic_Telecom_Packages.xlsx"
SHEETS = {"Packages": "packages", "Add_ons": "addons", "Customers": "customers"}


def seed_data(dataset_path: Path | str = DEFAULT_DATASET_PATH,
              database_path: Path | str = DEFAULT_DATABASE_PATH, *, reset: bool = False) -> dict[str, int]:
    """Upsert all seed records atomically; reset also removes stale catalog rows.

    Request/order history, checkpoints, and existing current-plan snapshots are
    never removed or overwritten. Reset only replaces original catalog rows.
    """
    workbook = load_workbook(dataset_path, read_only=True, data_only=True)
    catalogs = {}
    try:
        for sheet, table in SHEETS.items():
            if sheet not in workbook.sheetnames:
                raise ValueError(f"Missing worksheet: {sheet}")
            rows = workbook[sheet].iter_rows(values_only=True)
            if tuple(next(rows, ())) != CATALOG_COLUMNS[table]:
                raise ValueError(f"Unexpected columns in {sheet}")
            records = [row for row in rows if any(value is not None for value in row)]
            ids = [row[0] for row in records]
            if not records or any(not isinstance(key, str) or not key.strip() for key in ids):
                raise ValueError(f"Empty catalog or invalid IDs in {sheet}")
            if len(ids) != len(set(ids)):
                raise ValueError(f"Duplicate IDs in {sheet}")
            catalogs[table] = records
    finally:
        workbook.close()

    database = SearchDatabase(database_path)
    connection = database._connect()
    foreign_keys_enabled = connection.execute("PRAGMA foreign_keys").fetchone()[0]
    try:
        if reset and foreign_keys_enabled:
            # Catalog rows are replaced transactionally with the same IDs. Turn
            # FK enforcement off before BEGIN because SQLite's deferred-FK mode
            # still rejects some parent-table delete sequences. Re-enable it
            # immediately after the commit/rollback.
            connection.execute("PRAGMA foreign_keys = OFF")
        with connection:
            for table, records in catalogs.items():
                columns = CATALOG_COLUMNS[table]
                if reset:
                    connection.execute(f'DELETE FROM "{table}"')
                names = ", ".join(f'"{column}"' for column in columns)
                placeholders = ", ".join("?" for _ in columns)
                updates = ", ".join(f'"{column}" = excluded."{column}"' for column in columns[1:])
                connection.executemany(
                    f'INSERT INTO "{table}" ({names}) VALUES ({placeholders}) '
                    f'ON CONFLICT("{columns[0]}") DO UPDATE SET {updates}', records,
                )
    finally:
        if reset and foreign_keys_enabled:
            connection.execute("PRAGMA foreign_keys = ON")
        connection.close()
    # The first SearchDatabase construction can precede catalog seeding, so
    # initialize any customer snapshots that could not yet be resolved.
    from backend.services.current_plan_service import CurrentPlanService
    CurrentPlanService(database).initialize_missing_plans()
    return {table: len(records) for table, records in catalogs.items()}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dataset", type=Path, default=DEFAULT_DATASET_PATH)
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE_PATH)
    parser.add_argument("--reset", action="store_true", help="Replace catalogs only; preserve plans and history")
    args = parser.parse_args()
    for table, count in seed_data(args.dataset, args.database, reset=args.reset).items():
        print(f"{table}: {count} imported")


if __name__ == "__main__":
    main()
