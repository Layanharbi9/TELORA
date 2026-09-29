import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from backend.database.db import SearchDatabase
from backend.database.seed_data import DEFAULT_DATASET_PATH, SHEETS, seed_data
from backend.schemas.requirements import NeedSelection, Requirements
from backend.schemas.search import CandidatePool, SearchStatus
from backend.schemas.request import Request
from backend.schemas.checkpoint import WorkflowCheckpoint


class CatalogDatabaseTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.path = Path(directory.name) / "telora.sqlite3"
        self.database = SearchDatabase(self.path)

    def test_all_excel_values_and_types_round_trip(self):
        self.assertEqual(seed_data(database_path=self.path), {"packages": 100, "addons": 12, "customers": 50})
        workbook = load_workbook(DEFAULT_DATASET_PATH, read_only=True, data_only=True)
        try:
            for sheet, table in SHEETS.items():
                rows = workbook[sheet].iter_rows(values_only=True)
                headers = next(rows)
                expected = [dict(zip(headers, row)) for row in rows]
                actual = getattr(self.database, f"get_{table}")()
                self.assertEqual(actual, expected)
                for source, stored in zip(expected, actual):
                    for column in headers:
                        self.assertIs(type(stored[column]), type(source[column]), f"{table}.{column}")
        finally:
            workbook.close()

    def test_repeat_seed_and_reset_preserve_workflow_data(self):
        seed_data(database_path=self.path)
        request = Request(request_id="REQ-PRESERVE", customer_id="CUST-01",
                          requirements=Requirements(NeedSelection("Data", "Medium"), None, None, 400))
        self.database.create_request_with_checkpoint(request)
        checkpoint = WorkflowCheckpoint(request_id=request.request_id, candidate_pool=CandidatePool(
            package_candidates=[], addon_candidates=[], search_status=SearchStatus.NO_RESULTS))
        self.database.update_checkpoint(request.request_id, candidate_pool=checkpoint.candidate_pool)
        seed_data(database_path=self.path)
        connection = sqlite3.connect(self.path)
        try:
            with connection:
                connection.execute("UPDATE packages SET package_name = ? WHERE package_id = ?", ("changed", "PKG-LITE"))
                connection.execute("INSERT INTO packages (package_id) VALUES (?)", ("EXTRA",))
            seed_data(database_path=self.path)
            self.assertEqual(len(self.database.get_packages()), 101)
            self.assertNotEqual(next(row for row in self.database.get_packages() if row["package_id"] == "PKG-LITE")["package_name"], "changed")
            seed_data(database_path=self.path, reset=True)
            self.assertEqual(len(self.database.get_packages()), 100)
            self.assertEqual(len(self.database.get_addons()), 12)
            self.assertEqual(len(self.database.get_customers()), 50)
            self.assertEqual(self.database.get_request(request.request_id), request)
            self.assertEqual(self.database.get_checkpoint(request.request_id), checkpoint)
        finally:
            connection.close()

    def test_failed_reset_rolls_back_all_catalogs(self):
        seed_data(database_path=self.path)
        connection = sqlite3.connect(self.path)
        try:
            with connection:
                connection.execute("UPDATE packages SET package_name = ? WHERE package_id = ?", ("keep on rollback", "PKG-LITE"))
                connection.execute("CREATE TRIGGER reject_addons BEFORE INSERT ON addons BEGIN SELECT RAISE(ABORT, 'test failure'); END")
            before = self.database.get_packages()
            with self.assertRaises(sqlite3.IntegrityError):
                seed_data(database_path=self.path, reset=True)
            self.assertEqual(self.database.get_packages(), before)
            self.assertEqual(len(self.database.get_addons()), 12)
            self.assertEqual(len(self.database.get_customers()), 50)
        finally:
            connection.close()

    def test_invalid_workbooks_leave_existing_catalog_unchanged(self):
        seed_data(database_path=self.path)
        before = self.database.get_packages()
        for defect in ("missing_sheet", "bad_header", "duplicate_id"):
            with self.subTest(defect=defect):
                workbook = load_workbook(DEFAULT_DATASET_PATH)
                try:
                    if defect == "missing_sheet":
                        del workbook["Customers"]
                    elif defect == "bad_header":
                        workbook["Customers"]["A1"] = "wrong"
                    else:
                        workbook["Customers"]["A3"] = workbook["Customers"]["A2"].value
                    path = self.path.parent / f"{defect}.xlsx"
                    workbook.save(path)
                finally:
                    workbook.close()
                with self.assertRaises(ValueError):
                    seed_data(path, self.path, reset=True)
                self.assertEqual(self.database.get_packages(), before)

    def test_quotes_are_stored_as_values(self):
        workbook = load_workbook(DEFAULT_DATASET_PATH)
        value = "quoted'); DROP TABLE requests; --"
        try:
            workbook["Packages"]["B2"] = value
            path = self.path.parent / "quoted.xlsx"
            workbook.save(path)
        finally:
            workbook.close()
        seed_data(path, self.path)
        self.assertEqual(self.database.get_packages()[0]["package_name"], value)
        self.assertIsNone(self.database.get_request("absent"))


if __name__ == "__main__":
    unittest.main()
