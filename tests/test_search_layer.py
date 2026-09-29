import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
from contextlib import closing

from pydantic import ValidationError

from backend.database.seed_data import seed_data

from langchain_core.messages import AIMessage, ToolMessage
from langchain_core.tools import BaseTool

from backend.agents.search_agent import SEARCH_AGENT_SYSTEM_PROMPT, SEARCH_TOOLS, _ground_candidate_pool, build_search_context, run_search_agent, plan_and_retrieve, resolve_assessment_selection
from backend.database.db import SearchDatabase
from backend.schemas.requirements import NeedSelection, Requirements
from backend.schemas.search import AddonCandidate, AddonSearchInput, AddonSearchOutput, AssessmentSelection, CandidatePool, PackageCandidate, PackageSearchInput, PackageSearchOutput, SearchNeed, SearchJudgeResult, SearchPlan, SearchStatus, SearchMetadata, SearchTask, search_needs_from_requirements
from backend.schemas.state import SearchTrace, TELORAState
from backend.services.llm_retry import InvalidStructuredOutput
from backend.tools.addon_search_tool import _search_addons_impl, addon_search
from backend.tools.search_tool import _search_packages_impl, package_search


DATASET = Path(__file__).resolve().parents[1] / "data" / "TELORA_Synthetic_Telecom_Packages.xlsx"


def requirements(*selections):
    values = list(selections) + [None] * (3 - len(selections))
    return Requirements(values[0], values[1], values[2], 400)


def tool_call(name, needs, budget=400, call_id=None):
    """A native LangChain ToolCall dict, as it appears on AIMessage.tool_calls."""
    return {"name": name, "args": {"needs": needs, "preferred_budget_sar": budget},
            "id": call_id or f"call-{name}", "type": "tool_call"}


def ai_message(tool_calls=None, content=""):
    return AIMessage(content=content, tool_calls=tool_calls or [])


def bind_tool_calls(model, responses):
    """Configure model.bind_tools(...) to return a mock replaying `responses` in order."""
    tool_model = Mock()
    tool_model.invoke.side_effect = list(responses)
    model.bind_tools.return_value = tool_model
    return tool_model


class SearchInputTests(unittest.TestCase):
    def test_primary_only_priority_and_level(self):
        needs = search_needs_from_requirements(requirements(NeedSelection("Data", "Medium")))
        self.assertEqual(needs[0].model_dump(mode="json"), {"need": "Data", "level": "Medium", "priority": 1})

    def test_all_three_priorities_are_preserved(self):
        needs = search_needs_from_requirements(requirements(
            NeedSelection("Data", "Low"), NeedSelection("Gaming", "High"), NeedSelection("Calls", "Medium")
        ))
        self.assertEqual([need.priority for need in needs], [1, 2, 3])
        self.assertEqual([need.level.value for need in needs], ["Low", "High", "Medium"])

    def test_agent_context_preserves_requirements_and_search_needs(self):
        context = build_search_context(requirements(NeedSelection("Data", "Medium")))
        self.assertEqual(context["user_requirements"]["primary_need"]["need"], "Data")
        self.assertEqual(context["search_needs"][0]["priority"], 1)

    def test_agent_has_dedicated_retrieval_prompt(self):
        self.assertIn("TELORA's Search Agent", SEARCH_AGENT_SYSTEM_PROMPT)
        self.assertIn("NEVER invent", SEARCH_AGENT_SYSTEM_PROMPT)

    def test_grounding_returns_canonical_tool_records(self):
        metadata = {"tool_name": "package_search", "query_summary": "Data", "records_considered": 1, "records_returned": 1}
        tool_record = {
            "package_id": "TEL-P001", "name": "Catalog Package", "price_sar": 100.0,
            "matched_needs": [{"need": "Data", "level": "High", "priority": 1}],
            "match_reasons": ["Data capability is present in the catalog record"],
            "capabilities": {"data_gb": 20},
        }
        pool = CandidatePool(
            package_candidates=[PackageCandidate(**{**tool_record, "name": "LLM formatting"})],
            addon_candidates=[], search_status=SearchStatus.SUCCESS,
        )
        grounded = _ground_candidate_pool(pool, {"package_search": {**metadata, "candidates": [tool_record]}})
        self.assertEqual(grounded.package_candidates[0].name, "Catalog Package")



class CandidatePoolBoundaryTests(unittest.TestCase):
    def test_pool_contains_only_business_fields(self):
        fields = {"package_candidates", "addon_candidates", "search_status", "search_notes", "assessment_status"}
        self.assertEqual(set(CandidatePool.model_json_schema()["properties"]), fields)
        pool = CandidatePool(package_candidates=[], addon_candidates=[], search_status="NO_RESULTS")
        self.assertEqual(set(pool.model_dump()), fields)
        self.assertEqual(pool.search_notes, [])
        self.assertIsNone(pool.assessment_status)
        for invalid in ({"search_status": "INVALID"}, {"package_candidates": [{}]}):
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                CandidatePool.model_validate({**pool.model_dump(), **invalid})

    def test_tool_schemas_keep_validated_metadata(self):
        for schema in (PackageSearchOutput, AddonSearchOutput):
            with self.subTest(schema=schema):
                metadata = SearchMetadata(tool_name="test", query_summary="Data", records_considered=2, records_returned=0)
                output = schema(status="NO_RESULTS", candidates=[], search_metadata=metadata)
                self.assertIsInstance(output.search_metadata, SearchMetadata)
                with self.assertRaises(ValidationError):
                    schema(status="NO_RESULTS", candidates=[])
                with self.assertRaises(ValidationError):
                    schema(status="NO_RESULTS", candidates=[], search_metadata={**metadata.model_dump(), "records_considered": -1})

    def test_ids_selection_builds_pool_and_preserves_trace(self):
        args = {"needs": [{"need": "Data", "level": "Medium", "priority": 1}], "preferred_budget_sar": 400}
        record = {
            "package_id": "TEL-P001", "name": "Canonical package", "price_sar": 100,
            "matched_needs": args["needs"], "match_reasons": ["Data capability is present"],
            "capabilities": {"data_gb": 20},
        }
        outputs = [
            PackageSearchOutput(status="SUCCESS", candidates=[record], search_metadata=SearchMetadata(
                tool_name="package_search", query_summary="Data", records_considered=100,
                records_returned=1, execution_time_ms=1.25)),
            AddonSearchOutput(status="ERROR", candidates=[], error="Test retrieval error", search_metadata=SearchMetadata(
                tool_name="addon_search", query_summary="add-on retrieval failed", records_considered=0,
                records_returned=0, execution_time_ms=0.5)),
        ]
        fake_tools = {name: Mock(name=name) for name in ("package_search", "addon_search")}
        for (name, tool), output in zip(fake_tools.items(), outputs):
            tool.name = name
            tool.invoke.return_value = output
        model = Mock()
        # A single native tool-calling round: the model decides to call both tools itself.
        tool_model = bind_tool_calls(model, [
            ai_message(tool_calls=[
                tool_call("package_search", args["needs"], call_id="call-1"),
                tool_call("addon_search", args["needs"], call_id="call-2"),
            ]),
            ai_message(content="Compare Data suitability"),
        ])
        assessor = Mock()
        judge = Mock()
        judge.invoke.return_value = SearchJudgeResult(verdict="SUFFICIENT", gaps=[], evaluation_summary="Adequate alternatives")
        model.with_structured_output.side_effect = lambda schema, **kwargs: {AssessmentSelection: assessor, SearchJudgeResult: judge}[schema]
        # Exercise the internal IDs-only response before canonical reconstruction.
        assessor.invoke.side_effect = lambda messages: AssessmentSelection.model_validate({
            "package_ids": [record["package_id"]],
            "addon_ids": [], "search_status": "ERROR", "search_notes": ["Add-on search failed"],
        })
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", fake_tools):
            result = run_search_agent(TELORAState("REQ-BOUNDARY", requirements(NeedSelection("Data", "Medium"))), llm=model, judge_llm=model)
        model.bind_tools.assert_called_once_with(SEARCH_TOOLS)
        self.assertEqual(tool_model.invoke.call_count, 2)
        self.assertEqual([call.args[0] for call in model.with_structured_output.call_args_list], [AssessmentSelection, SearchJudgeResult])
        for tool in fake_tools.values():
            tool.invoke.assert_called_once_with(args)
        # Both tool results returned as ToolMessages to the SAME conversation before the model continued.
        second_round_messages = tool_model.invoke.call_args_list[1].args[0]
        tool_messages = [message for message in second_round_messages if isinstance(message, ToolMessage)]
        self.assertEqual({message.tool_call_id for message in tool_messages}, {"call-1", "call-2"})
        self.assertEqual(json.loads(tool_messages[0].content)["status"], "SUCCESS")
        expected_plan = SearchPlan(tasks=[
            {"tool_name": "package_search", "need_priorities": [1]},
            {"tool_name": "addon_search", "need_priorities": [1]},
        ], assessment_focus="Compare Data suitability")
        self.assertEqual(result.search_trace.search_plan, expected_plan)
        self.assertEqual([stage.stage for stage in result.search_trace.stages],
                         ["PLAN", "RETRIEVE", "REDUCE", "ASSESS", "LOOKUP", "JUDGE", "OUTPUT", "ROUTE"])
        self.assertEqual(result.search_trace.stages[1].status, "ERROR")
        for stage in [stage for stage in result.search_trace.stages[1:] if stage.stage not in ("JUDGE", "ROUTE")]:
            self.assertEqual(stage.details["package_ids"], [record["package_id"]])
            self.assertEqual(stage.details["addon_ids"], [])
        self.assertEqual(result.search_trace.stages[3].details["search_notes"], ["Add-on search failed"])
        output_stage = next(s for s in result.search_trace.stages if s.stage == "OUTPUT")
        self.assertEqual(output_stage.details["grounding"], "SUCCESS")
        # The addon_search tool itself errored, so routing is deterministically
        # ERROR regardless of the assessor's own "ERROR" claim happening to agree.
        route_stage = result.search_trace.stages[-1]
        self.assertEqual(route_stage.stage, "ROUTE")
        self.assertEqual(route_stage.details["routing_status"], "ERROR")
        self.assertEqual(route_stage.details["routing_reason"], "TOOL_ERROR")
        self.assertEqual(route_stage.details["assessment_status"], "ERROR")
        self.assertEqual(route_stage.details["package_count"], 1)
        self.assertEqual(route_stage.details["addon_count"], 0)
        self.assertEqual(result.candidate_pool.search_status.value, "ERROR")
        self.assertEqual(result.candidate_pool.assessment_status.value, "ERROR")
        self.assertEqual([trace.task_index for trace in result.search_trace.tool_traces], [1, 2])
        self.assertEqual(result.candidate_pool.package_candidates[0].name, record["name"])
        self.assertNotIn("search_metadata", result.candidate_pool.model_dump())
        for trace, output in zip(result.search_trace.tool_traces, outputs):
            self.assertEqual(trace.output, output.model_dump(mode="json"))
            self.assertEqual(trace.input, args)
            self.assertGreaterEqual(trace.execution_time_ms, 0)
        messages = assessor.invoke.call_args.args[0]
        context = json.loads(messages[-1].content)
        self.assertEqual(context["search_plan"], expected_plan.model_dump(mode="json"))
        self.assertEqual(context["user_requirements"], result.user_requirements.to_dict())
        self.assertEqual(context["tool_outputs"], {name: output.model_dump(mode="json") for name, output in zip(fake_tools, outputs)})


class ModelRoutingTests(unittest.TestCase):
    def test_independent_model_configuration_and_defaults(self):
        from backend.agents.search_agent import create_search_llm, create_judge_llm
        with patch.dict("os.environ", {}, clear=True), patch("backend.agents.search_agent.ChatOpenAI") as client:
            create_search_llm()
            self.assertEqual(client.call_args.kwargs, {"model": "gpt-4o-mini", "temperature": 0, "max_retries": 0})
            create_judge_llm()
            self.assertEqual(client.call_args.kwargs, {"model": "gpt-5.6-terra", "temperature": 0, "max_retries": 0, "reasoning_effort": "none"})
        for judge_name in ("judge-a", "judge-b"):
            with patch.dict("os.environ", {"OPENAI_MODEL": "search-custom", "JUDGE_MODEL": judge_name}), patch("backend.agents.search_agent.ChatOpenAI") as client:
                create_search_llm()
                self.assertEqual(client.call_args.kwargs["model"], "search-custom")
                create_judge_llm()
                self.assertEqual(client.call_args.kwargs["model"], judge_name)

    def test_separate_clients_route_all_stages(self):
        for verdict in ("SUFFICIENT", "INSUFFICIENT"):
            with self.subTest(verdict=verdict):
                search, judge, assessor, judge_response = [Mock() for _ in range(4)]
                req = requirements(NeedSelection("Data", "Medium"))
                plan = SearchPlan(tasks=[{"tool_name": "package_search", "need_priorities": [1]}], assessment_focus="Fit")
                assessor.invoke.return_value = AssessmentSelection(package_ids=[], addon_ids=[], search_status="NO_RESULTS")
                search.with_structured_output.side_effect = lambda schema, **kwargs: {AssessmentSelection: assessor}[schema]
                judge.with_structured_output.return_value = judge_response
                judge_response.invoke.return_value = SearchJudgeResult(verdict=verdict, gaps=[] if verdict == "SUFFICIENT" else [
                    {"need_priority": 1, "gap_type": "WEAK_COVERAGE", "reason": "Coverage gap"}], evaluation_summary="Summary")
                with patch("backend.agents.search_agent.create_search_llm", return_value=search) as search_factory, \
                     patch("backend.agents.search_agent.create_judge_llm", return_value=judge) as judge_factory, \
                     patch("backend.agents.search_agent.plan_and_retrieve", return_value=(plan, {}, [])) as planned:
                    run_search_agent(TELORAState("ROUTING", req))
                search_factory.assert_called_once_with()
                judge_factory.assert_called_once_with()
                planned.assert_called_once_with(req, search, 6)
                judge.with_structured_output.assert_called_once_with(SearchJudgeResult, method="function_calling")
                expected = [AssessmentSelection] + ([AssessmentSelection] if verdict == "INSUFFICIENT" else [])
                self.assertEqual([call.args[0] for call in search.with_structured_output.call_args_list], expected)
                self.assertEqual(assessor.invoke.call_count, 2 if verdict == "INSUFFICIENT" else 1)


class RuntimeJudgeTests(unittest.TestCase):
    def test_invalid_judge_exhausts_retry_and_unknown_id_fails_explicitly(self):
        req = requirements(NeedSelection("Data", "Medium"))
        needs_arg = [{"need": "Data", "level": "Medium", "priority": 1}]
        for failure in ("priority", "schema", "unknown_id"):
            with self.subTest(failure=failure):
                model, assessor, judge, tool = [Mock() for _ in range(4)]
                bind_tool_calls(model, [
                    ai_message(tool_calls=[tool_call("package_search", needs_arg)]),
                    ai_message(content="Fit"),
                ])
                initial = AssessmentSelection(package_ids=[], addon_ids=[], search_status="NO_RESULTS")
                invalid_reassessment = AssessmentSelection(package_ids=["UNKNOWN"], addon_ids=[], search_status="SUCCESS")
                assessor.invoke.side_effect = [initial, invalid_reassessment, invalid_reassessment]
                judgment = {"verdict": "INSUFFICIENT", "gaps": [{"need_priority": 2 if failure == "priority" else 1,
                    "gap_type": "WEAK_COVERAGE", "reason": "Missing coverage"}], "evaluation_summary": "Needs improvement"}
                if failure == "schema":
                    judgment["gaps"] = []
                judge.invoke.return_value = judgment
                model.with_structured_output.side_effect = lambda schema, **kwargs: {AssessmentSelection: assessor, SearchJudgeResult: judge}[schema]
                tool.invoke.return_value = PackageSearchOutput(status="NO_RESULTS", candidates=[], search_metadata=SearchMetadata(
                    tool_name="package_search", query_summary="Data", records_considered=0, records_returned=0))
                with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool}), self.assertRaises(ValueError):
                    run_search_agent(TELORAState("FAIL", req), llm=model, judge_llm=model)
                tool.invoke.assert_called_once()
                self.assertEqual(judge.invoke.call_count, 1 if failure == "unknown_id" else 2)
                # Canonical ID lookup is outside technical LLM retry; the invalid
                # reassessment is rejected after its single semantic call.
                self.assertEqual(assessor.invoke.call_count, 2 if failure == "unknown_id" else 1)

    def test_judge_contract_and_priority_validation(self):
        gap = {"need_priority": 1, "gap_type": "WEAK_COVERAGE", "reason": "Weak fit"}
        for verdict, gaps in (("SUFFICIENT", [gap]), ("INSUFFICIENT", [])):
            with self.subTest(verdict=verdict), self.assertRaises(ValidationError):
                SearchJudgeResult(verdict=verdict, gaps=gaps, evaluation_summary="Summary")
        for invalid in ({**gap, "need_priority": True}, {**gap, "gap_type": "OTHER"}, {**gap, "reason": " "}):
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                SearchJudgeResult(verdict="INSUFFICIENT", gaps=[invalid], evaluation_summary="Summary")
        result = SearchJudgeResult(verdict="INSUFFICIENT", gaps=[{**gap, "need_priority": 2}], evaluation_summary="Summary")
        with self.assertRaises(ValueError):
            result.validate_requirements(requirements(NeedSelection("Data", "Medium")))

    def test_bounded_reassessment_all_final_statuses_and_trace(self):
        # Routing is deterministic (packages exist, no tool error -> SUCCESS)
        # regardless of what the assessor/reassessor itself claims; that raw
        # claim is preserved separately as assessment_status, advisory only.
        req = requirements(NeedSelection("Data", "Medium"))
        needs_arg = [{"need": "Data", "level": "Medium", "priority": 1}]
        records = [{"package_id": key, "name": key, "price_sar": 100,
                    "matched_needs": [], "match_reasons": [], "capabilities": {"data_gb": 20}}
                   for key in ("A", "B", "C")]
        output = PackageSearchOutput(status="SUCCESS", candidates=records, search_metadata=SearchMetadata(
            tool_name="package_search", query_summary="Data", records_considered=3, records_returned=3))
        for status in ("SUCCESS", "INSUFFICIENT", "ERROR"):
            with self.subTest(status=status):
                first = AssessmentSelection(package_ids=["A"], addon_ids=[], search_status="SUCCESS", search_notes=["Initial note"])
                final_ids = ["C", "B"]
                second = AssessmentSelection(package_ids=final_ids, addon_ids=[], search_status=status, search_notes=["Refined note"])
                judgment = SearchJudgeResult(verdict="INSUFFICIENT", gaps=[
                    {"need_priority": 1, "gap_type": kind, "reason": "Need distinct relevant alternatives"}
                    for kind in ("WEAK_COVERAGE", "LIMITED_OPTIONS")], evaluation_summary="Insufficient diversity potential")
                assessor, judge, model, tool = Mock(), Mock(), Mock(), Mock()
                bind_tool_calls(model, [
                    ai_message(tool_calls=[tool_call("package_search", needs_arg)]),
                    ai_message(content="Fit"),
                ])
                assessor.invoke.side_effect = [first, second]
                judge.invoke.return_value = judgment
                model.with_structured_output.side_effect = lambda schema, **kwargs: {AssessmentSelection: assessor, SearchJudgeResult: judge}[schema]
                tool.invoke.return_value = output
                with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool}):
                    result = run_search_agent(TELORAState("RUN", req), llm=model, judge_llm=model)
                tool.invoke.assert_called_once()
                judge.invoke.assert_called_once()
                self.assertEqual(assessor.invoke.call_count, 2)
                self.assertEqual([call.args[0] for call in model.with_structured_output.call_args_list],
                                 [AssessmentSelection, SearchJudgeResult, AssessmentSelection])
                judge_context = json.loads(judge.invoke.call_args.args[0][-1].content)
                self.assertEqual(judge_context["shortlist"]["package_candidates"][0]["capabilities"], records[0]["capabilities"])
                context = json.loads(assessor.invoke.call_args_list[1].args[0][-1].content)
                self.assertEqual(context["user_requirements"], req.to_dict())
                self.assertEqual(context["tool_outputs"]["package_search"], output.model_dump(mode="json"))
                self.assertEqual(context["initial_selection"], first.model_dump(mode="json"))
                self.assertEqual(context["judge_result"], judgment.model_dump(mode="json"))
                self.assertEqual([p.package_id for p in result.candidate_pool.package_candidates], final_ids)
                # Routing status is always SUCCESS here: >=1 final package, no tool error.
                self.assertEqual(result.candidate_pool.search_status.value, "SUCCESS")
                self.assertEqual(result.candidate_pool.assessment_status.value, status)
                stages = {stage.stage: stage for stage in result.search_trace.stages}
                self.assertTrue(stages["JUDGE"].details["refinement_triggered"])
                self.assertEqual(stages["JUDGE"].details["evaluation_summary"], judgment.evaluation_summary)
                self.assertEqual(stages["ASSESS"].details["search_notes"], ["Initial note"])
                self.assertEqual(stages["REASSESS"].details["search_notes"], ["Refined note"])
                self.assertEqual(stages["REASSESS"].details["changes"]["package_ids"]["after"], final_ids)
                self.assertEqual(stages["ROUTE"].details["routing_status"], "SUCCESS")
                self.assertEqual(stages["ROUTE"].details["routing_reason"], "PACKAGES_AVAILABLE")
                self.assertEqual(stages["ROUTE"].details["assessment_status"], status)
                self.assertEqual(stages["ROUTE"].details["package_count"], 2)

    def test_reassess_selecting_zero_packages_despite_retrieval_is_invalid_and_recovers_on_retry(self):
        # Retrieval found packages; REASSESS drops every one of them. That must
        # be treated as invalid output (not genuine NO_RESULTS) and retried --
        # the existing bounded technical retry, not a new mechanism.
        req = requirements(NeedSelection("Data", "Medium"))
        needs_arg = [{"need": "Data", "level": "Medium", "priority": 1}]
        records = [{"package_id": key, "name": key, "price_sar": 100,
                    "matched_needs": [], "match_reasons": [], "capabilities": {"data_gb": 20}}
                   for key in ("A", "B")]
        output = PackageSearchOutput(status="SUCCESS", candidates=records, search_metadata=SearchMetadata(
            tool_name="package_search", query_summary="Data", records_considered=2, records_returned=2))
        first = AssessmentSelection(package_ids=["A"], addon_ids=[], search_status="SUCCESS")
        empty_reassessment = AssessmentSelection(package_ids=[], addon_ids=[], search_status="NO_RESULTS")
        recovered = AssessmentSelection(package_ids=["B"], addon_ids=[], search_status="SUCCESS")
        judgment = SearchJudgeResult(verdict="INSUFFICIENT",
                                     gaps=[{"need_priority": 1, "gap_type": "WEAK_COVERAGE", "reason": "Gap"}],
                                     evaluation_summary="Gap")
        assessor, judge, model, tool = Mock(), Mock(), Mock(), Mock()
        bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("package_search", needs_arg)]),
            ai_message(content="Fit"),
        ])
        assessor.invoke.side_effect = [first, empty_reassessment, recovered]
        judge.invoke.return_value = judgment
        model.with_structured_output.side_effect = lambda schema, **kwargs: {AssessmentSelection: assessor, SearchJudgeResult: judge}[schema]
        tool.invoke.return_value = output
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool}):
            result = run_search_agent(TELORAState("RUN", req), llm=model, judge_llm=model)
        self.assertEqual(assessor.invoke.call_count, 3)
        self.assertEqual([p.package_id for p in result.candidate_pool.package_candidates], ["B"])
        self.assertEqual(result.candidate_pool.search_status.value, "SUCCESS")

    def test_reassess_selecting_zero_packages_exhausts_retry_and_raises(self):
        # Same as above, but the retry also drops every package -> ERROR, not
        # a false "your needs are impossible" NO_RESULTS.
        req = requirements(NeedSelection("Data", "Medium"))
        needs_arg = [{"need": "Data", "level": "Medium", "priority": 1}]
        records = [{"package_id": "A", "name": "A", "price_sar": 100,
                    "matched_needs": [], "match_reasons": [], "capabilities": {"data_gb": 20}}]
        output = PackageSearchOutput(status="SUCCESS", candidates=records, search_metadata=SearchMetadata(
            tool_name="package_search", query_summary="Data", records_considered=1, records_returned=1))
        first = AssessmentSelection(package_ids=["A"], addon_ids=[], search_status="SUCCESS")
        empty_reassessment = AssessmentSelection(package_ids=[], addon_ids=[], search_status="NO_RESULTS")
        judgment = SearchJudgeResult(verdict="INSUFFICIENT",
                                     gaps=[{"need_priority": 1, "gap_type": "WEAK_COVERAGE", "reason": "Gap"}],
                                     evaluation_summary="Gap")
        assessor, judge, model, tool = Mock(), Mock(), Mock(), Mock()
        bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("package_search", needs_arg)]),
            ai_message(content="Fit"),
        ])
        assessor.invoke.side_effect = [first, empty_reassessment, empty_reassessment]
        judge.invoke.return_value = judgment
        model.with_structured_output.side_effect = lambda schema, **kwargs: {AssessmentSelection: assessor, SearchJudgeResult: judge}[schema]
        tool.invoke.return_value = output
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool}), \
                self.assertLogs("backend.services.llm_retry", level="ERROR"), \
                self.assertRaises(InvalidStructuredOutput):
            run_search_agent(TELORAState("RUN", req), llm=model, judge_llm=model)
        self.assertEqual(assessor.invoke.call_count, 3)


class SearchProgressCallbackTests(unittest.TestCase):
    """on_phase marks the real FINDING -> CHECKING boundary and nothing else."""

    def setUp(self):
        self.req = requirements(NeedSelection("Data", "Medium"))
        self.needs_arg = [{"need": "Data", "level": "Medium", "priority": 1}]
        # A is cheaper, B has more data: a real trade-off, so reduce keeps both.
        records = [{"package_id": key, "name": key, "price_sar": price, "matched_needs": self.needs_arg,
                    "match_reasons": [], "capabilities": {"data_gb": data}}
                   for key, price, data in (("A", 80, 20), ("B", 100, 40))]
        self.output = PackageSearchOutput(status="SUCCESS", candidates=records, search_metadata=SearchMetadata(
            tool_name="package_search", query_summary="Data", records_considered=2, records_returned=2))

    def run_agent(self, on_phase=None, events=None):
        model, assessor, judge, tool = Mock(), Mock(), Mock(), Mock()
        bind_tool_calls(model, [ai_message(tool_calls=[tool_call("package_search", self.needs_arg)]),
                                ai_message(content="Fit")])

        def assess(messages):
            if events is not None:
                events.append("ASSESS")
            return AssessmentSelection(package_ids=["B", "A"], addon_ids=[], search_status="SUCCESS")

        def retrieve(args):
            if events is not None:
                events.append("RETRIEVE")
            return self.output

        assessor.invoke.side_effect = assess
        judge.invoke.return_value = SearchJudgeResult(verdict="SUFFICIENT", gaps=[], evaluation_summary="Fine")
        model.with_structured_output.side_effect = lambda schema, **kwargs: {
            AssessmentSelection: assessor, SearchJudgeResult: judge}[schema]
        tool.invoke.side_effect = retrieve
        kwargs = {} if on_phase is None else {"on_phase": on_phase}
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool}):
            return run_search_agent(TELORAState("RUN", self.req), llm=model, judge_llm=model, **kwargs)

    def test_fires_exactly_once_after_reduce_and_before_assess(self):
        events = []
        result = self.run_agent(on_phase=lambda phase: events.append(phase), events=events)
        self.assertEqual(events, ["RETRIEVE", "CHECKING", "ASSESS"])
        # The deterministic reduce already ran (its stage exists) when CHECKING began.
        self.assertIn("REDUCE", [stage.stage for stage in result.search_trace.stages])

    def test_does_not_fire_when_search_fails_before_the_boundary(self):
        calls = []
        model = Mock()
        bad_needs = [{"need": "Gaming", "level": "High", "priority": 1}]
        bind_tool_calls(model, [ai_message(tool_calls=[tool_call("package_search", bad_needs)]),
                                ai_message(tool_calls=[tool_call("package_search", bad_needs)])])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": Mock()}), \
             self.assertLogs("backend.services.llm_retry", level="ERROR"), \
             self.assertRaises(InvalidStructuredOutput):
            run_search_agent(TELORAState("RUN", self.req), llm=model, judge_llm=model, on_phase=calls.append)
        self.assertEqual(calls, [])

    def test_search_result_is_identical_with_and_without_the_callback(self):
        without = self.run_agent()
        with_callback = self.run_agent(on_phase=lambda phase: None)
        self.assertEqual(with_callback.candidate_pool, without.candidate_pool)
        self.assertEqual(with_callback.search_trace.search_plan, without.search_trace.search_plan)
        strip = lambda trace: [(s.stage, s.status, {k: v for k, v in s.details.items() if k != "llm_retry"})
                               for s in trace.stages]
        self.assertEqual(strip(with_callback.search_trace), strip(without.search_trace))


class SelectionLookupTests(unittest.TestCase):
    def test_selection_schema_rejects_objects_duplicates_and_invalid_status(self):
        base = dict(package_ids=[], addon_ids=[], search_status="SUCCESS")
        self.assertEqual(set(AssessmentSelection.model_fields), {"package_ids", "addon_ids", "search_status", "search_notes"})
        for invalid in ({"package_ids": [{}]}, {"package_ids": ["A", "A"]},
                        {"addon_ids": ["X", "X"]}, {"package_ids": [""]},
                        {"search_status": "invalid"}, {"package_candidates": []}):
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                AssessmentSelection(**{**base, **invalid})

    def test_lookup_and_grounding_preserve_exact_order_and_omissions(self):
        packages = [{"package_id": key, "name": key, "price_sar": price,
                     "matched_needs": [], "match_reasons": ["Canonical"], "capabilities": {"data_gb": 20}}
                    for key, price in (("A", 10), ("B", 20), ("C", 900))]
        addons = [{"addon_id": key, "name": key, "price_sar": 5,
                   "matched_needs": [], "match_reasons": [], "capabilities": {}} for key in ("X", "Y", "Z")]
        outputs = {"package_search": {"candidates": packages}, "addon_search": {"candidates": addons}}
        selection = AssessmentSelection(package_ids=["C", "A"], addon_ids=["Z", "X"], search_status="SUCCESS", search_notes=["Tradeoff"])
        pool = resolve_assessment_selection(selection, outputs)
        grounded = _ground_candidate_pool(pool, outputs)
        self.assertEqual(grounded, pool)
        self.assertEqual([p.model_dump() for p in pool.package_candidates], [PackageCandidate(**packages[2]).model_dump(), PackageCandidate(**packages[0]).model_dump()])
        self.assertEqual([a.addon_id for a in pool.addon_candidates], ["Z", "X"])
        self.assertEqual(pool.search_notes, selection.search_notes)
        empty = resolve_assessment_selection(AssessmentSelection(package_ids=[], addon_ids=[], search_status="NO_RESULTS"), outputs)
        self.assertEqual(empty.package_candidates, [])
        self.assertEqual(empty.addon_candidates, [])
        for field in ("package_ids", "addon_ids"):
            bad = AssessmentSelection(**{**selection.model_dump(), field: ["unknown"]})
            with self.subTest(field=field), self.assertRaisesRegex(ValueError, "Unknown selected"):
                resolve_assessment_selection(bad, outputs)
        selection.package_ids.append("C")
        with self.assertRaisesRegex(ValueError, "Duplicate selected"):
            resolve_assessment_selection(selection, outputs)


class SearchPlanningTests(unittest.TestCase):
    def test_invalid_plan_structure_is_rejected(self):
        valid = {"tool_name": "package_search", "need_priorities": [1]}
        for tasks in ([], [valid, valid], [{**valid, "tool_name": "calculation"}],
                      [{**valid, "need_priorities": []}], [{**valid, "need_priorities": [1, 1]}],
                      [{**valid, "need_priorities": [0]}], [{**valid, "need_priorities": [True]}],
                      [{**valid, "preferred_budget_sar": 1}]):
            with self.subTest(tasks=tasks), self.assertRaises(ValidationError):
                SearchPlan(tasks=tasks, assessment_focus="Compare needs")
        with self.assertRaises(ValidationError):
            SearchPlan(tasks=[valid], assessment_focus=" ")

    def test_tools_execute_in_call_order_with_original_arguments(self):
        """The model drives retrieval via native tool calls across rounds; Python
        executes each call, in the order requested, with the original grounded args."""
        req = Requirements(NeedSelection("Data", "Medium"), NeedSelection("Streaming", "High"), NeedSelection("Calls", "Low"), 220)
        original = [need.model_dump(mode="json") for need in search_needs_from_requirements(req)]
        model = Mock()
        bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("addon_search", [original[1]], budget=220, call_id="call-1")]),
            ai_message(tool_calls=[tool_call("package_search", [original[2], original[0]], budget=220, call_id="call-2")]),
            ai_message(content="Compare fit"),
        ])
        calls = []
        tools = {}
        for name, schema in (("addon_search", AddonSearchOutput), ("package_search", PackageSearchOutput)):
            output = schema(status="NO_RESULTS", candidates=[], search_metadata=SearchMetadata(
                tool_name=name, query_summary="none", records_considered=0, records_returned=0))
            tools[name] = Mock()
            def invoke(args, name=name, output=output):
                calls.append((name, args))
                return output
            tools[name].invoke.side_effect = invoke
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", tools):
            plan, outputs, traces = plan_and_retrieve(req, model, max_tool_rounds=6)
        self.assertEqual([name for name, _ in calls], ["addon_search", "package_search"])
        self.assertEqual(calls[0][1], {"needs": [original[1]], "preferred_budget_sar": 220})
        self.assertEqual(calls[1][1], {"needs": [original[2], original[0]], "preferred_budget_sar": 220})
        self.assertEqual([trace.task_index for trace in traces], [1, 2])
        self.assertEqual(plan.tasks, [SearchTask(tool_name="addon_search", need_priorities=[2]),
                                       SearchTask(tool_name="package_search", need_priorities=[1, 3])])
        self.assertEqual(plan.assessment_focus, "Compare fit")

    def test_grounding_preserves_subset_order_and_rejects_unknown_ids(self):
        records = [{"package_id": key, "name": key, "price_sar": 10,
                    "matched_needs": [], "match_reasons": [], "capabilities": {}}
                   for key in ("A", "B", "C")]
        pool = CandidatePool(package_candidates=[records[2], records[0]], addon_candidates=[], search_status="SUCCESS")
        outputs = {"package_search": {"candidates": records}}
        self.assertEqual([c.package_id for c in _ground_candidate_pool(pool, outputs).package_candidates], ["C", "A"])
        pool.package_candidates[0].package_id = "unknown"
        with self.assertRaises(ValueError):
            _ground_candidate_pool(pool, outputs)

    def test_tool_trace_defaults_and_serialization(self):
        from backend.schemas.state import ToolTrace
        trace = SearchTrace("llm_search_agent", [ToolTrace("package_search", {}, {}, "NO_RESULTS", 1)])
        self.assertIsNone(trace.search_plan)
        self.assertEqual(trace.stages, [])
        self.assertEqual(trace.to_dict()["tool_traces"][0]["status"], "NO_RESULTS")


class NativeToolCallGroundingTests(unittest.TestCase):
    """Exercises plan_and_retrieve directly: the LLM decides which Search Tool(s)
    to call and with which needs; Python only grounds, executes, and bounds rounds."""

    def setUp(self):
        self.req = requirements(NeedSelection("Data", "Medium"), NeedSelection("Streaming", "High"))
        self.canonical = [need.model_dump(mode="json") for need in search_needs_from_requirements(self.req)]

    def fake_tool(self, name="package_search"):
        tool = Mock()
        tool.invoke.return_value = PackageSearchOutput(status="NO_RESULTS", candidates=[], search_metadata=SearchMetadata(
            tool_name=name, query_summary="Data", records_considered=0, records_returned=0))
        return tool

    def test_valid_native_package_search_call_retrieves_and_completes(self):
        tool = self.fake_tool()
        model = Mock()
        bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("package_search", self.canonical)]),
            ai_message(content="Compare"),
        ])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool, "addon_search": Mock()}):
            plan, outputs, traces = plan_and_retrieve(self.req, model, max_tool_rounds=6)
        tool.invoke.assert_called_once()
        self.assertEqual(plan.tasks, [SearchTask(tool_name="package_search", need_priorities=[1, 2])])
        self.assertEqual(plan.assessment_focus, "Compare")
        self.assertIn("package_search", outputs)
        self.assertEqual([trace.tool_name for trace in traces], ["package_search"])

    def test_valid_native_addon_search_call_retrieves_and_completes(self):
        tool = self.fake_tool("addon_search")
        model = Mock()
        bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("addon_search", self.canonical)]),
            ai_message(content="Compare"),
        ])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": Mock(), "addon_search": tool}):
            plan, outputs, traces = plan_and_retrieve(self.req, model, max_tool_rounds=6)
        tool.invoke.assert_called_once()
        self.assertEqual(plan.tasks, [SearchTask(tool_name="addon_search", need_priorities=[1, 2])])
        self.assertEqual([trace.tool_name for trace in traces], ["addon_search"])

    def test_multiple_native_tool_calls_across_rounds(self):
        package_tool, addon_tool = self.fake_tool(), self.fake_tool("addon_search")
        model = Mock()
        bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("package_search", [self.canonical[0]], call_id="call-1")]),
            ai_message(tool_calls=[tool_call("addon_search", [self.canonical[1]], call_id="call-2")]),
            ai_message(content="Compare"),
        ])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": package_tool, "addon_search": addon_tool}):
            plan, outputs, traces = plan_and_retrieve(self.req, model, max_tool_rounds=6)
        package_tool.invoke.assert_called_once()
        addon_tool.invoke.assert_called_once()
        self.assertEqual([trace.tool_name for trace in traces], ["package_search", "addon_search"])
        self.assertEqual(set(outputs), {"package_search", "addon_search"})

    def test_tool_results_return_as_toolmessage_to_the_same_conversation(self):
        tool = self.fake_tool()
        model = Mock()
        tool_model = bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("package_search", self.canonical, call_id="call-xyz")]),
            ai_message(content="Compare"),
        ])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool, "addon_search": Mock()}):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        second_call_messages = tool_model.invoke.call_args_list[1].args[0]
        tool_messages = [message for message in second_call_messages if isinstance(message, ToolMessage)]
        self.assertEqual(len(tool_messages), 1)
        self.assertEqual(tool_messages[0].tool_call_id, "call-xyz")
        self.assertEqual(json.loads(tool_messages[0].content)["status"], "NO_RESULTS")
        # The AIMessage that requested the call is carried forward in the same conversation.
        self.assertTrue(any(isinstance(message, AIMessage) and message.tool_calls for message in second_call_messages))

    def test_invented_need_is_rejected_before_execution(self):
        tool = self.fake_tool()
        model = Mock()
        bad_needs = [{"need": "Gaming", "level": "High", "priority": 1}]
        bind_tool_calls(model, [ai_message(tool_calls=[tool_call("package_search", bad_needs)])])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool, "addon_search": Mock()}), \
             self.assertRaises(InvalidStructuredOutput):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        tool.invoke.assert_not_called()

    def test_altered_level_is_rejected_before_execution(self):
        tool = self.fake_tool()
        model = Mock()
        altered = [{**self.canonical[0], "level": "High"}]
        bind_tool_calls(model, [ai_message(tool_calls=[tool_call("package_search", altered)])])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool, "addon_search": Mock()}), \
             self.assertRaises(InvalidStructuredOutput):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        tool.invoke.assert_not_called()

    def test_altered_priority_is_rejected_before_execution(self):
        tool = self.fake_tool()
        model = Mock()
        altered = [{**self.canonical[0], "priority": 2}]
        bind_tool_calls(model, [ai_message(tool_calls=[tool_call("package_search", altered)])])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool, "addon_search": Mock()}), \
             self.assertRaises(InvalidStructuredOutput):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        tool.invoke.assert_not_called()

    def test_altered_budget_is_rejected_before_execution(self):
        tool = self.fake_tool()
        model = Mock()
        bind_tool_calls(model, [ai_message(tool_calls=[tool_call("package_search", self.canonical, budget=999)])])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool, "addon_search": Mock()}), \
             self.assertRaises(InvalidStructuredOutput):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        tool.invoke.assert_not_called()

    def test_incomplete_coverage_is_rejected_once_the_model_stops(self):
        """A partial, individually-valid call still executes; coverage is only
        knowable -- and only checked -- once the model stops calling tools."""
        tool = self.fake_tool()
        model = Mock()
        bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("package_search", [self.canonical[0]])]),
            ai_message(content="Compare"),
        ])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": tool, "addon_search": Mock()}), \
             self.assertRaises(InvalidStructuredOutput):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        tool.invoke.assert_called_once()

    def test_zero_tool_calls_is_rejected_as_insufficient_coverage(self):
        model = Mock()
        bind_tool_calls(model, [ai_message(content="Nothing to search")])
        tools = {"package_search": Mock(), "addon_search": Mock()}
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", tools), self.assertRaises(ValidationError):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        for tool in tools.values():
            tool.invoke.assert_not_called()

    def test_duplicate_tool_call_across_rounds_is_rejected(self):
        package_tool = self.fake_tool()
        model = Mock()
        bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("package_search", [self.canonical[0]], call_id="call-1")]),
            ai_message(tool_calls=[tool_call("package_search", [self.canonical[1]], call_id="call-2")]),
        ])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": package_tool, "addon_search": Mock()}), \
             self.assertRaises(InvalidStructuredOutput):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        package_tool.invoke.assert_called_once()  # only the first, valid call executes

    def test_excess_call_within_a_single_round_is_rejected(self):
        package_tool = self.fake_tool()
        model = Mock()
        bind_tool_calls(model, [ai_message(tool_calls=[
            tool_call("package_search", [self.canonical[0]], call_id="call-1"),
            tool_call("package_search", [self.canonical[1]], call_id="call-2"),
        ])])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": package_tool, "addon_search": Mock()}), \
             self.assertRaises(InvalidStructuredOutput):
            plan_and_retrieve(self.req, model, max_tool_rounds=6)
        package_tool.invoke.assert_called_once()

    def test_max_tool_rounds_bounds_the_loop(self):
        package_tool, addon_tool = self.fake_tool(), self.fake_tool("addon_search")
        model = Mock()
        tool_model = bind_tool_calls(model, [
            ai_message(tool_calls=[tool_call("package_search", [self.canonical[0]])]),
            ai_message(tool_calls=[tool_call("addon_search", [self.canonical[1]])]),
        ])
        with patch("backend.agents.search_agent.SEARCH_TOOL_MAP", {"package_search": package_tool, "addon_search": addon_tool}), \
             self.assertRaises(InvalidStructuredOutput):
            plan_and_retrieve(self.req, model, max_tool_rounds=1)
        self.assertEqual(tool_model.invoke.call_count, 1)
        package_tool.invoke.assert_called_once()
        addon_tool.invoke.assert_not_called()


class NativeToolTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.database_path = Path(directory.name) / "catalog.sqlite3"
        seed_data(DATASET, self.database_path)
        # Public tool calls use the isolated SQLite catalog, never the user's DB.
        for module in ("backend.tools.search_tool", "backend.tools.addon_search_tool"):
            patcher = patch(f"{module}.SearchDatabase", side_effect=lambda path: SearchDatabase(self.database_path))
            patcher.start()
            self.addCleanup(patcher.stop)

    def test_runtime_tools_do_not_open_excel(self):
        with patch("openpyxl.load_workbook", side_effect=AssertionError("Runtime Excel access")), patch("zipfile.ZipFile", side_effect=AssertionError("Runtime workbook access")):
            package = package_search.invoke({"needs": [{"need": "Data", "level": "Medium", "priority": 1}], "preferred_budget_sar": 400})
            addon = addon_search.invoke({"needs": [{"need": "Gaming", "level": "High", "priority": 1}], "preferred_budget_sar": 400})
        self.assertEqual(package.status, SearchStatus.SUCCESS)
        self.assertEqual(addon.status, SearchStatus.SUCCESS)

    def test_runtime_reads_updated_sqlite_values(self):
        database = SearchDatabase(self.database_path)
        from contextlib import closing
        with closing(database._connect()) as connection, connection:
            connection.execute("UPDATE packages SET package_name = ?", ("SQLite package",))
            connection.execute("UPDATE addons SET addon_name = ?", ("SQLite addon",))
        for tool, need, expected in ((package_search, "Data", "SQLite package"), (addon_search, "Gaming", "SQLite addon")):
            output = tool.invoke({"needs": [{"need": need, "level": "Medium", "priority": 1}], "preferred_budget_sar": 400})
            self.assertEqual(output.status, SearchStatus.SUCCESS)
            self.assertTrue(output.candidates)
            self.assertTrue(all(candidate.name == expected for candidate in output.candidates))

    def test_exports_are_real_framework_tools(self):
        self.assertIsInstance(package_search, BaseTool)
        self.assertIsInstance(addon_search, BaseTool)
        self.assertEqual(package_search.args_schema, PackageSearchInput)
        self.assertEqual(addon_search.args_schema, AddonSearchInput)
        self.assertEqual([tool.name for tool in SEARCH_TOOLS], ["package_search", "addon_search"])

    def test_package_tool_returns_grounded_packages(self):
        query = PackageSearchInput(needs=[SearchNeed(need="Data", level="Medium", priority=1)], preferred_budget_sar=400)
        output = package_search.invoke(query.model_dump())
        self.assertEqual(output.status, SearchStatus.SUCCESS)
        self.assertTrue(output.candidates)
        import openpyxl
        workbook = openpyxl.load_workbook(DATASET, read_only=True, data_only=True)
        rows = workbook["Packages"].iter_rows(values_only=True)
        headers = next(rows)
        records = {row[headers.index("package_id")]: row[headers.index("package_name")] for row in rows}
        workbook.close()
        self.assertTrue({candidate.package_id for candidate in output.candidates} <= set(records))
        self.assertTrue(all(records[candidate.package_id] == candidate.name for candidate in output.candidates))

    def test_addon_tool_returns_only_grounded_addons(self):
        query = AddonSearchInput(needs=[SearchNeed(need="Gaming", level="High", priority=1)], preferred_budget_sar=400)
        output = addon_search.invoke(query.model_dump())
        self.assertIn(output.status, {SearchStatus.SUCCESS, SearchStatus.NO_RESULTS})
        import openpyxl
        workbook = openpyxl.load_workbook(DATASET, read_only=True, data_only=True)
        rows = workbook["Add_ons"].iter_rows(values_only=True)
        headers = next(rows)
        records = {row[headers.index("addon_id")]: row[headers.index("addon_name")] for row in rows}
        workbook.close()
        self.assertTrue({candidate.addon_id for candidate in output.candidates} <= set(records))
        self.assertTrue(all(records[candidate.addon_id] == candidate.name for candidate in output.candidates))

    def test_empty_and_error_results_are_structured(self):
        query = PackageSearchInput(needs=[SearchNeed(need="International Calls", level="Low", priority=1)], preferred_budget_sar=400)
        with patch("backend.tools.search_tool.SearchDatabase", SearchDatabase):
            output = _search_packages_impl(query, self.database_path.parent / "missing.sqlite3")
        self.assertEqual(output.status, SearchStatus.ERROR)
        self.assertTrue(output.error)
        self.assertEqual(output.candidates, [])
        self.assertIsNotNone(output.search_metadata.execution_time_ms)

        addon_output = _search_addons_impl(
            AddonSearchInput(needs=[SearchNeed(need="Data Sharing", level="Low", priority=1)], preferred_budget_sar=400),
            self.database_path,
        )
        self.assertIn(addon_output.status, {SearchStatus.SUCCESS, SearchStatus.NO_RESULTS})


if __name__ == "__main__":
    unittest.main()
